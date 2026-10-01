"""Case Study 5 report app: export of a reviewed draft to Word (.docx) and Markdown."""

from __future__ import annotations

import io
from datetime import datetime, timezone

import report_core as rc

SEVERITY_ORDER = {"error": 0, "warning": 1, "info": 2}


def _sorted(findings: list[rc.Finding]) -> list[rc.Finding]:
    return sorted(findings, key=lambda f: (SEVERITY_ORDER.get(f.severity, 3), f.check))


def to_markdown(company: rc.Company, year: int, slots: dict[str, str], method: str, findings: list[rc.Finding],
                signoff: dict) -> str:
    out = [f"# {company.name}: Solvency and Financial Condition Report {year}", "",
           f"## Sections E.1 and E.2 (draft: {method})", "",
           f"*{company.text.get('fictitious_note', 'Fictitious insurer; synthetic figures.')}*", ""]
    for section, title in (("E.1", "E.1 Own funds"), ("E.2", "E.2 Solvency Capital Requirement and Minimum Capital "
                                                              "Requirement")):
        out += [f"### {title}", ""]
        for slot in rc.SLOT_IDS:
            if rc.SLOT_SECTION[slot] != section:
                continue
            out += [f"**{rc.SLOT_TITLE[slot]}**", ""]
            out += [p.strip() for p in (slots.get(slot, "") or "").split("\n") if p.strip()]
            out.append("")
            if slot == "e1_structure":
                out += _md_table("Table E.1: Own funds by tier (EUR million)", rc.table_e1(company, year))
            if slot == "e2_amounts":
                out += _md_table("Table E.2: SCR by risk module, MCR and coverage ratios (EUR million)",
                                 rc.table_e2(company, year))
    out += ["---", "", "## Checks", ""]
    if findings:
        out += ["| Check | Severity | Where | Finding |", "|---|---|---|---|"]
        for f in _sorted(findings):
            out.append(f"| {f.check} | {f.severity} | {f.slot or '(document)'} | {f.message.replace('|', '/')} |")
    else:
        out.append("No findings.")
    out += ["", "## Sign-off", "", f"- Reviewer: {signoff.get('reviewer') or '(not signed)'}",
            f"- Decision: {signoff.get('decision') or '(none)'}", f"- Notes: {signoff.get('notes') or '-'}",
            f"- Exported: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}", ""]
    return "\n".join(out)


def _md_table(caption: str, rows: list[dict]) -> list[str]:
    if not rows:
        return []
    cols = list(rows[0])
    lines = [f"*{caption}*", "", "| " + " | ".join(c or " " for c in cols) + " |",
             "|" + "|".join("---" if i == 0 else "---:" for i in range(len(cols))) + "|"]
    for r in rows:
        lines.append("| " + " | ".join(_cell(r[c]) for c in cols) + " |")
    return lines + [""]


def _cell(value) -> str:
    if isinstance(value, float):
        return f"{value:,.1f}"
    return str(value)


def to_docx(company: rc.Company, year: int, slots: dict[str, str], method: str, findings: list[rc.Finding],
            signoff: dict) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    navy = RGBColor(0x0F, 0x2B, 0x72)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.color.rgb = navy
        doc.styles[name].font.name = "Arial"

    doc.add_heading(f"{company.name}: Solvency and Financial Condition Report {year}", level=0)
    p = doc.add_paragraph()
    run = p.add_run(f"Sections E.1 and E.2 - draft: {method}. "
                    + company.text.get("fictitious_note", "Fictitious insurer; synthetic figures."))
    run.italic = True

    def table(caption: str, rows: list[dict]) -> None:
        if not rows:
            return
        cap = doc.add_paragraph()
        cap.add_run(caption).italic = True
        cols = list(rows[0])
        t = doc.add_table(rows=1, cols=len(cols))
        t.style = "Light Grid Accent 1"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, c in enumerate(cols):
            t.rows[0].cells[i].text = c
        for r in rows:
            cells = t.add_row().cells
            for i, c in enumerate(cols):
                cells[i].text = _cell(r[c])
                if i:
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        doc.add_paragraph()

    for section, title in (("E.1", "E.1 Own funds"),
                           ("E.2", "E.2 Solvency Capital Requirement and Minimum Capital Requirement")):
        doc.add_heading(title, level=1)
        for slot in rc.SLOT_IDS:
            if rc.SLOT_SECTION[slot] != section:
                continue
            doc.add_heading(rc.SLOT_TITLE[slot], level=3)
            for para in (slots.get(slot, "") or "").split("\n"):
                if para.strip():
                    doc.add_paragraph(para.strip())
            if slot == "e1_structure":
                table("Table E.1: Own funds by tier (EUR million)", rc.table_e1(company, year))
            if slot == "e2_amounts":
                table("Table E.2: SCR by risk module, MCR and coverage ratios (EUR million)",
                      rc.table_e2(company, year))

    doc.add_page_break()
    doc.add_heading("Checks", level=1)
    if findings:
        t = doc.add_table(rows=1, cols=4)
        t.style = "Light Grid Accent 1"
        for i, c in enumerate(("Check", "Severity", "Where", "Finding")):
            t.rows[0].cells[i].text = c
        for f in _sorted(findings):
            cells = t.add_row().cells
            cells[0].text, cells[1].text = f.check, f.severity
            cells[2].text, cells[3].text = f.slot or "(document)", f.message
    else:
        doc.add_paragraph("No findings.")
    doc.add_heading("Sign-off", level=1)
    for label, key in (("Reviewer", "reviewer"), ("Decision", "decision"), ("Notes", "notes")):
        doc.add_paragraph(f"{label}: {signoff.get(key) or '-'}")
    doc.add_paragraph(f"Exported: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
