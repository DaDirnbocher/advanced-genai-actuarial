"""Case Study 5 report app: the workbook, the figures, the placeholders, the old method and the checks.

Nothing in this module calls a language model or Streamlit. It holds:
- the metrics of SFCR sections E.1 (own funds) and E.2 (SCR and MCR) and how code derives totals, the amounts
  eligible under the tier limits and the coverage ratios from the inputs of the workbook;
- the workbook reader (one Excel file per company: figures 2021-2026, driver bullets, the texts published for
  2022-2025, the house phrases and the terminology);
- the facts catalogue of a reporting year: every figure as a {{placeholder}} with its rendered value;
- the old method: last year's published text with the figures updated by code;
- the deterministic checks that every draft goes through, whichever method wrote it.
"""

from __future__ import annotations

import difflib
import io
import re
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any

import openpyxl

# ---------------------------------------------------------------------------
# Metrics
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Metric:
    id: str
    label: str
    unit: str                    # "eur_m" (EUR million) or "pct" (a ratio in per cent)
    kind: str                    # "input" (typed in the workbook) or "derived" (computed by code)
    section: str                 # "E.1", "E.2" or "E.1(e)" (the reconciliation of equity)
    synonyms: tuple[str, ...] = ()
    dp: int = 1                  # decimals shown in the text
    movement: bool = True        # whether the text describes its change with a movement word
    signed: bool = False         # a reconciliation item: its sign says whether it adds or deducts


METRICS: list[Metric] = [
    # E.2: the SCR by risk module (standard formula). Deductions are stored as positive amounts.
    Metric("scr_market", "market risk", "eur_m", "input", "E.2", ("market risk",)),
    Metric("scr_default", "counterparty default risk", "eur_m", "input", "E.2",
           ("counterparty default risk", "counterparty risk", "default risk")),
    Metric("scr_life", "life underwriting risk", "eur_m", "input", "E.2", ("life underwriting risk", "life risk")),
    Metric("scr_health", "health underwriting risk", "eur_m", "input", "E.2",
           ("health underwriting risk", "health risk")),
    Metric("scr_nonlife", "non-life underwriting risk", "eur_m", "input", "E.2",
           ("non-life underwriting risk", "non-life risk")),
    Metric("div_benefit", "diversification benefit (deducted)", "eur_m", "input", "E.2",
           ("diversification benefit", "diversification")),
    Metric("bscr", "Basic Solvency Capital Requirement (BSCR)", "eur_m", "derived", "E.2",
           ("basic solvency capital requirement", "bscr")),
    Metric("scr_operational", "operational risk", "eur_m", "input", "E.2", ("operational risk",)),
    Metric("lac_tp", "loss-absorbing capacity of technical provisions (deducted)", "eur_m", "input", "E.2",
           ("loss-absorbing capacity of technical provisions",)),
    Metric("lac_dt", "loss-absorbing capacity of deferred taxes (deducted)", "eur_m", "input", "E.2",
           ("loss-absorbing capacity of deferred taxes",)),
    Metric("scr", "Solvency Capital Requirement (SCR)", "eur_m", "derived", "E.2",
           ("solvency capital requirement", "scr")),
    Metric("mcr", "Minimum Capital Requirement (MCR)", "eur_m", "input", "E.2",
           ("minimum capital requirement", "mcr")),
    # E.1: own funds by tier, as available and as eligible under the tier limits.
    Metric("of_t1", "Tier 1 own funds", "eur_m", "input", "E.1", ("tier 1",)),
    Metric("of_t2", "Tier 2 own funds", "eur_m", "input", "E.1", ("tier 2",)),
    Metric("of_t3", "Tier 3 own funds", "eur_m", "input", "E.1", ("tier 3",)),
    Metric("of_total", "total basic own funds", "eur_m", "derived", "E.1", ("basic own funds", "total own funds")),
    Metric("eof_scr_t1", "Tier 1 eligible to cover the SCR", "eur_m", "derived", "E.1", (), movement=False),
    Metric("eof_scr_t2", "Tier 2 eligible to cover the SCR", "eur_m", "derived", "E.1", (), movement=False),
    Metric("eof_scr_t3", "Tier 3 eligible to cover the SCR", "eur_m", "derived", "E.1", (), movement=False),
    Metric("eof_scr", "eligible own funds to cover the SCR", "eur_m", "derived", "E.1",
           ("eligible own funds to cover the scr", "own funds eligible to cover the scr",
            "eligible own funds to cover the solvency capital requirement",
            "own funds eligible to cover the solvency capital requirement", "eligible own funds")),
    Metric("eof_mcr_t1", "Tier 1 eligible to cover the MCR", "eur_m", "derived", "E.1", (), movement=False),
    Metric("eof_mcr_t2", "Tier 2 eligible to cover the MCR", "eur_m", "derived", "E.1", (), movement=False),
    Metric("eof_mcr", "eligible own funds to cover the MCR", "eur_m", "derived", "E.1",
           ("eligible own funds to cover the mcr", "own funds eligible to cover the mcr",
            "eligible own funds to cover the minimum capital requirement",
            "own funds eligible to cover the minimum capital requirement")),
    Metric("of_inelig_scr", "own funds not eligible to cover the SCR (tier limits)", "eur_m", "derived", "E.1",
           (), movement=False),
    Metric("cov_scr", "SCR coverage ratio", "pct", "derived", "E.2",
           ("scr coverage ratio", "solvency ii ratio", "scr coverage", "solvency ratio", "coverage of the scr"), dp=0),
    Metric("cov_mcr", "MCR coverage ratio", "pct", "derived", "E.2",
           ("mcr coverage ratio", "mcr coverage", "coverage of the mcr"), dp=0),
    # E.1(e): equity in the financial statements -> excess of assets over liabilities -> basic own funds.
    Metric("equity_fs", "equity in the financial statements", "eur_m", "input", "E.1(e)",
           ("equity in the financial statements", "equity in our financial statements",
            "equity in its financial statements", "equity under ifrs"), movement=False),
    Metric("rec_tp", "valuation of technical provisions (difference)", "eur_m", "input", "E.1(e)",
           ("technical provisions",), movement=False, signed=True),
    Metric("rec_intangibles", "intangible assets (difference)", "eur_m", "input", "E.1(e)",
           ("intangible assets",), movement=False, signed=True),
    Metric("rec_dt", "deferred taxes (difference)", "eur_m", "input", "E.1(e)",
           ("deferred taxes", "deferred tax"), movement=False, signed=True),
    Metric("rec_other", "other valuation differences", "eur_m", "input", "E.1(e)",
           ("other valuation differences", "other differences"), movement=False, signed=True),
    Metric("eaol", "excess of assets over liabilities", "eur_m", "derived", "E.1(e)",
           ("excess of assets over liabilities",), movement=False),
    Metric("foreseeable_div", "foreseeable dividends (deducted)", "eur_m", "input", "E.1",
           ("foreseeable dividend",), movement=False),
    Metric("sub_liab", "subordinated liabilities in basic own funds", "eur_m", "input", "E.1",
           (), movement=False),
]
METRIC: dict[str, Metric] = {m.id: m for m in METRICS}
INPUT_IDS = [m.id for m in METRICS if m.kind == "input"]
MODULES = ["scr_market", "scr_default", "scr_life", "scr_health", "scr_nonlife"]
RATIOS = {"cov_scr", "cov_mcr"}

MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December"]


def _d(x: float) -> Decimal:
    return Decimal(str(x))


def q1(x: Decimal | float) -> float:
    """Round to EUR 0.1 million, half away from zero, as Excel's ROUND does."""
    return float(Decimal(str(x)).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP))


def derive(inputs: dict[str, float]) -> dict[str, float]:
    """All figures of one year-end from the inputs: BSCR and SCR, the own funds by tier as available and as eligible
    under the tier limits of Art. 82(1) of Delegated Regulation (EU) 2015/35, the coverage ratios and the excess of
    assets over liabilities. Code derives every total and ratio; nothing derived is typed."""
    v = {k: float(inputs.get(k, 0.0) or 0.0) for k in INPUT_IDS}
    bscr = sum(_d(v[m]) for m in MODULES) - _d(v["div_benefit"])
    scr = bscr + _d(v["scr_operational"]) - _d(v["lac_tp"]) - _d(v["lac_dt"])
    t1, t2, t3 = _d(v["of_t1"]), _d(v["of_t2"]), _d(v["of_t3"])
    mcr = _d(v["mcr"])
    # SCR: tier 3 below 15% of the SCR; tier 2 plus tier 3 at most 50% of the SCR (tier 1 is at least 50%).
    e_t3 = min(t3, Decimal("0.15") * scr)
    e_t2 = max(Decimal(0), min(t2, Decimal("0.5") * scr - e_t3))
    # MCR: tier 2 at most 20% of the MCR; tier 3 is not eligible.
    m_t2 = min(t2, Decimal("0.2") * mcr)
    out = dict(v)
    out.update(
        bscr=q1(bscr), scr=q1(scr), of_total=q1(t1 + t2 + t3),
        eof_scr_t1=q1(t1), eof_scr_t2=q1(e_t2), eof_scr_t3=q1(e_t3),
        eof_mcr_t1=q1(t1), eof_mcr_t2=q1(m_t2),
    )
    out["eof_scr"] = q1(_d(out["eof_scr_t1"]) + _d(out["eof_scr_t2"]) + _d(out["eof_scr_t3"]))
    out["eof_mcr"] = q1(_d(out["eof_mcr_t1"]) + _d(out["eof_mcr_t2"]))
    out["of_inelig_scr"] = q1(_d(out["of_total"]) - _d(out["eof_scr"]))
    out["cov_scr"] = 100.0 * out["eof_scr"] / out["scr"] if out["scr"] else 0.0
    out["cov_mcr"] = 100.0 * out["eof_mcr"] / out["mcr"] if out["mcr"] else 0.0
    out["eaol"] = q1(_d(v["equity_fs"]) + _d(v["rec_tp"]) + _d(v["rec_intangibles"]) + _d(v["rec_dt"])
                     + _d(v["rec_other"]))
    return out


def consistency_findings(figures: dict[int, dict[str, float]]) -> list[dict]:
    """Checks of the workbook's figures themselves, year by year: the limits and identities that must hold."""
    rows = []
    for year, f in figures.items():
        def add(severity: str, message: str) -> None:
            rows.append({"year": year, "severity": severity, "message": message})

        if f["scr"] <= 0 or f["mcr"] <= 0:
            add("error", "the SCR and the MCR must be positive")
            continue
        if f["of_t1"] < 0.5 * f["scr"]:
            add("error", "Tier 1 is below half of the SCR: the tier limits of Art. 82(1) are not met")
        if f["of_t1"] < 0.8 * f["mcr"]:
            add("error", "Tier 1 is below 80% of the MCR: the tier limits of Art. 82(2) are not met")
        if not 0.25 * f["scr"] - 0.05 <= f["mcr"] <= 0.45 * f["scr"] + 0.05:
            add("warning", "the MCR lies outside the corridor of 25% to 45% of the SCR")
        if f["eof_scr"] < f["of_total"] - 0.05:
            add("info", f"tier limits bind: EUR {f['of_total'] - f['eof_scr']:,.1f} million of own funds are not "
                        "eligible to cover the SCR")
        expected = q1(_d(f["eaol"]) - _d(f["foreseeable_div"]) + _d(f["sub_liab"]))
        if abs(expected - f["of_total"]) > 0.05:
            add("error", f"basic own funds (EUR {f['of_total']:,.1f} million) differ from the excess of assets over "
                         f"liabilities minus foreseeable dividends plus subordinated liabilities "
                         f"(EUR {expected:,.1f} million)")
        if f["cov_scr"] < 100:
            add("error", "the SCR is not covered")
    return rows


# ---------------------------------------------------------------------------
# Number formats (as in the Case Study 5 notebook)
# ---------------------------------------------------------------------------

def fmt_eur(value: float) -> str:
    return f"EUR {abs(value):,.1f} million"


def fmt_pct(value: float, dp: int = 0) -> str:
    return f"{value:,.{dp}f}%"


def fmt_pp(value: float) -> str:
    n = round(abs(value))
    return f"{n} percentage point" + ("" if n == 1 else "s")


def fmt_date(year: int) -> str:
    return f"31 December {year}"


# ---------------------------------------------------------------------------
# The workbook
# ---------------------------------------------------------------------------

SHEETS = ("Company", "Figures", "Bullets", "Narratives", "House phrases", "Terminology")


@dataclass
class Company:
    """One company's workbook, read into plain Python."""
    company_id: str
    name: str
    text: dict[str, str]                         # Company sheet: field -> value
    years: list[int]                             # columns of the Figures sheet
    inputs: dict[int, dict[str, float]]          # year -> input metric -> value
    figures: dict[int, dict[str, float]]         # year -> every metric, derived by code
    bullets: dict[int, list[dict]]               # year -> [{id, topic, text}]
    narratives: dict[int, list[dict]]            # year -> [{section, slot, title, based_on, text}]
    phrases: list[dict]                          # house phrases: applies_to, from, to, up, down, stable
    terminology: list[dict]                      # preferred, avoid (list), note
    workbook_cached: dict[int, dict[str, float]] = field(default_factory=dict)  # derived values cached by Excel
    source: str = ""

    @property
    def current_year(self) -> int:
        """The year to draft: the last year with figures and bullets but no published text."""
        candidates = [y for y in self.years if y in self.bullets and y not in self.narratives]
        return max(candidates) if candidates else max(self.years)

    @property
    def history_years(self) -> list[int]:
        return sorted(y for y in self.narratives if y < self.current_year)

    def slots(self, year: int) -> dict[str, str]:
        return {row["slot"]: row["text"] for row in self.narratives.get(year, [])}


class WorkbookError(ValueError):
    """The workbook cannot be read: a sheet, a column or a value is missing or malformed."""


def _rows(ws) -> list[dict[str, Any]]:
    values = list(ws.iter_rows(values_only=True))
    if not values:
        return []
    header = [str(h).strip() if h is not None else "" for h in values[0]]
    out = []
    for raw in values[1:]:
        if raw is None or all(c is None or str(c).strip() == "" for c in raw):
            continue
        out.append({header[i]: raw[i] for i in range(min(len(header), len(raw))) if header[i]})
    return out


def _num(value: Any, where: str) -> float:
    if value is None or (isinstance(value, str) and not value.strip()):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError as exc:
        raise WorkbookError(f"{where}: {value!r} is not a number") from exc


def load_workbook(source: str | Path | bytes, source_name: str = "") -> Company:
    """Read a company workbook. Only the input rows of the Figures sheet are used; code derives the rest."""
    if isinstance(source, (bytes, bytearray)):
        wb = openpyxl.load_workbook(io.BytesIO(source), data_only=True, read_only=True)
        source_name = source_name or "uploaded workbook"
    else:
        wb = openpyxl.load_workbook(source, data_only=True, read_only=True)
        source_name = source_name or Path(source).name
    missing = [s for s in SHEETS if s not in wb.sheetnames]
    if missing:
        raise WorkbookError(f"sheet(s) missing: {', '.join(missing)}")

    text = {str(r.get("field", "")).strip(): str(r.get("value", "") or "").strip() for r in _rows(wb["Company"])}
    if not text.get("company_name"):
        raise WorkbookError("Company sheet: the field company_name is empty")

    fig_rows = _rows(wb["Figures"])
    if not fig_rows:
        raise WorkbookError("Figures sheet is empty")
    years = sorted(int(k) for k in fig_rows[0] if re.fullmatch(r"\d{4}", str(k).strip()))
    if len(years) < 2:
        raise WorkbookError("Figures sheet: at least two year columns (e.g. 2025 and 2026) are needed")
    by_id = {str(r.get("id", "")).strip(): r for r in fig_rows}
    inputs: dict[int, dict[str, float]] = {y: {} for y in years}
    cached: dict[int, dict[str, float]] = {y: {} for y in years}
    for mid in INPUT_IDS:
        row = by_id.get(mid)
        for y in years:
            inputs[y][mid] = _num(row.get(str(y), row.get(y)) if row else 0.0, f"Figures, {mid}, {y}")
    for m in METRICS:
        if m.kind == "derived" and m.id in by_id:
            for y in years:
                val = by_id[m.id].get(str(y), by_id[m.id].get(y))
                if isinstance(val, (int, float)):
                    cached[y][m.id] = float(val)
    figures = {y: derive(inputs[y]) for y in years}

    bullets: dict[int, list[dict]] = {}
    for r in _rows(wb["Bullets"]):
        y = int(_num(r.get("year"), "Bullets, year"))
        bullets.setdefault(y, []).append({"id": str(r.get("id", "")).strip(), "topic": str(r.get("topic", "") or "").strip(),
                                          "text": str(r.get("text", "") or "").strip()})
    narratives: dict[int, list[dict]] = {}
    for r in _rows(wb["Narratives"]):
        y = int(_num(r.get("year"), "Narratives, year"))
        narratives.setdefault(y, []).append({
            "section": str(r.get("section", "")).strip(), "slot": str(r.get("slot", "")).strip(),
            "title": str(r.get("title", "") or "").strip(), "based_on": str(r.get("based_on", "") or "").strip(),
            "text": str(r.get("text", "") or "").strip()})
    phrases = []
    for r in _rows(wb["House phrases"]):
        phrases.append({"applies_to": str(r.get("applies_to", "")).strip().lower(),
                        "from": _num(r.get("from"), "House phrases, from"),
                        "to": _num(r.get("to"), "House phrases, to") if r.get("to") not in (None, "") else 1e9,
                        "up": str(r.get("phrase_up", "") or "").strip(),
                        "down": str(r.get("phrase_down", "") or "").strip()})
    terminology = []
    for r in _rows(wb["Terminology"]):
        avoid = [a.strip() for a in str(r.get("avoid", "") or "").split(";") if a.strip()]
        terminology.append({"preferred": str(r.get("preferred", "") or "").strip(), "avoid": avoid,
                            "note": str(r.get("note", "") or "").strip()})
    wb.close()
    return Company(company_id=text.get("company_id", "uploaded"), name=text["company_name"], text=text, years=years,
                   inputs=inputs, figures=figures, bullets=bullets, narratives=narratives, phrases=phrases,
                   terminology=terminology, workbook_cached=cached, source=source_name)


def cached_differences(company: Company) -> list[dict]:
    """Derived values that Excel cached in the workbook but that differ from the values code derives."""
    rows = []
    for y, values in company.workbook_cached.items():
        for mid, cached in values.items():
            mine = company.figures[y].get(mid)
            dp = METRIC[mid].dp
            if mine is not None and abs(round(cached, dp + 1) - round(mine, dp + 1)) > 0.5 * 10 ** -dp:
                rows.append({"year": y, "metric": METRIC[mid].label, "workbook": cached, "code": mine})
    return rows


# ---------------------------------------------------------------------------
# House phrases
# ---------------------------------------------------------------------------

def house_phrase(company: Company, metric_id: str, change: float) -> str | None:
    """The house phrase for a change: per cent for amounts, percentage points for ratios."""
    kind = "ratio" if metric_id in RATIOS else "amount"
    size = abs(change)
    for p in company.phrases:
        if p["applies_to"] == kind and p["from"] <= size < p["to"]:
            return p["up"] if change > 0 else p["down"]
    return None


# ---------------------------------------------------------------------------
# The facts catalogue of a reporting year
# ---------------------------------------------------------------------------

@dataclass
class Fact:
    id: str
    metric: str          # metric id, or "year" / "date"
    period: str          # "cur", "prev" or "change"
    kind: str            # "level", "change_pct", "change_eur", "change_pp", "year", "date"
    value: float
    display: str
    label: str
    direction: str = ""  # for changes: "increase", "decrease" or "unchanged"
    phrase: str = ""     # for changes: the house phrase


def facts_for(company: Company, year: int) -> dict[str, Fact]:
    """Every figure of a reporting year as a placeholder: levels at the two year-ends, changes, years and dates."""
    if year - 1 not in company.figures:
        raise ValueError(f"no figures for {year - 1}, the comparison year of {year}")
    cur, prev = company.figures[year], company.figures[year - 1]
    facts: dict[str, Fact] = {}

    def add(f: Fact) -> None:
        facts[f.id] = f

    add(Fact("year_cur", "year", "cur", "year", year, str(year), "reporting year"))
    add(Fact("year_prev", "year", "prev", "year", year - 1, str(year - 1), "previous year"))
    add(Fact("date_cur", "date", "cur", "date", year, fmt_date(year), "end of the reporting year"))
    add(Fact("date_prev", "date", "prev", "date", year - 1, fmt_date(year - 1), "end of the previous year"))
    for k in range(2, 6):
        add(Fact(f"year_m{k}", "year", "older", "year", year - k, str(year - k),
                 f"{k} years before the reporting year (for events dated in that year)"))
    for cid, value, label in constants(company):
        add(Fact(cid, "const", "const", "const", value, fmt_pct(value), label))
    for m in METRICS:
        c, p = cur[m.id], prev[m.id]
        if abs(c) < 1e-9 and abs(p) < 1e-9:
            continue  # an item this company does not have (e.g. life underwriting risk at a non-life insurer)
        for period, val in (("cur", c), ("prev", p)):
            display = fmt_pct(val, m.dp) if m.unit == "pct" else fmt_eur(val)
            add(Fact(f"{m.id}_{period}", m.id, period, "level", val, display, m.label))
        if not m.movement:
            continue
        if m.unit == "pct":
            pp = c - p
            direction = "unchanged" if round(abs(pp)) == 0 else ("increase" if pp > 0 else "decrease")
            add(Fact(f"{m.id}_change_pp", m.id, "change", "change_pp", pp, fmt_pp(pp), m.label, direction,
                     house_phrase(company, m.id, pp) or ""))
        else:
            diff = q1(_d(c) - _d(p))
            direction = "unchanged" if abs(diff) < 0.05 else ("increase" if diff > 0 else "decrease")
            add(Fact(f"{m.id}_change_eur", m.id, "change", "change_eur", diff, fmt_eur(diff), m.label, direction))
            if abs(p) > 1e-9:
                pct = 100.0 * (c / p - 1.0)
                direction_pct = "unchanged" if round(abs(pct), 1) == 0 else ("increase" if pct > 0 else "decrease")
                add(Fact(f"{m.id}_change_pct", m.id, "change", "change_pct", pct, fmt_pct(abs(pct), 1), m.label,
                         direction_pct, house_phrase(company, m.id, pct) or ""))
    return facts


REGULATORY_CONSTANTS = [
    ("lim_t3_scr", 15.0, "limit: Tier 3 below this share of the SCR (Art. 82(1))"),
    ("lim_t2t3_scr", 50.0, "limit: Tier 2 plus Tier 3 at most this share of the SCR (Art. 82(1))"),
    ("lim_t2_mcr", 20.0, "limit: Tier 2 at most this share of the MCR (Art. 82(2))"),
    ("mcr_floor", 25.0, "floor of the MCR as a share of the SCR (Art. 129(3) of the Directive)"),
    ("mcr_cap", 45.0, "cap of the MCR as a share of the SCR (Art. 129(3) of the Directive)"),
]


def constants(company: Company) -> list[tuple[str, float, str]]:
    """Percentages that the text may quote but that are no figures of the year: the risk appetite from the Company
    sheet and the regulatory limits. They are placeholders too, so that code owns these numbers as well."""
    out = []
    for key, label in (("appetite_low", "risk appetite: SCR coverage ratio of at least / lower end of the range"),
                       ("appetite_high", "risk appetite: upper end of the range")):
        raw = str(company.text.get(key, "") or "").strip().rstrip("%")
        if raw:
            try:
                out.append((key, float(raw), label))
            except ValueError:
                pass
    return out + REGULATORY_CONSTANTS


def movements(facts: dict[str, Fact]) -> dict[str, Fact]:
    """The change of each metric that has one: per cent for amounts (if defined), percentage points for ratios."""
    out: dict[str, Fact] = {}
    for f in facts.values():
        if f.kind in ("change_pct", "change_pp"):
            out[f.metric] = f
        elif f.kind == "change_eur" and f.metric not in out and f"{f.metric}_change_pct" not in facts:
            out[f.metric] = f
    return out


# ---------------------------------------------------------------------------
# Placeholders: render, and recover from a published text
# ---------------------------------------------------------------------------

PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}")


def render(template: str, facts: dict[str, Fact]) -> str:
    def fill(m: re.Match) -> str:
        f = facts.get(m.group(1))
        return f.display if f else f"[unknown {m.group(1)}]"
    return PLACEHOLDER.sub(fill, template)


# Numbers as the house formats print them, longest patterns first.
NUM_EUR = r"EUR\s?(?P<eur>-?\d[\d,]*(?:\.\d+)?)\s?million"
NUM_PP = r"(?P<pp>\d+(?:\.\d+)?)\s?percentage points?"
NUM_PCT = r"(?P<pct>-?\d+(?:\.\d+)?)\s?%"
NUM_DATE = r"(?P<day>\d{1,2})\s(?P<month>" + "|".join(MONTHS) + r")\s(?P<dyear>\d{4})"
NUM_YEAR = r"\b(?P<year>(?:19|20)\d\d)\b"
NUMBER = re.compile("|".join([NUM_EUR, NUM_PP, NUM_PCT, NUM_DATE, NUM_YEAR]))
# Digits that are names, not figures: Tier 1/2/3, Table E.1/E.2, Article 82 and the like.
REFERENCE = re.compile(r"\b(?:Tier\s?[123]|Table\s?E\.[12]|E\.[1-6]|Art(?:icle|\.)\s?\d+[a-z]?(?:\(\d+\))*)")
STRAY_DIGITS = re.compile(r"\d[\d,]*(?:\.\d+)?")


def _value_of(m: re.Match) -> tuple[str, float, int]:
    """(kind, value, decimals printed) of a number match."""
    for kind in ("eur", "pp", "pct"):
        if m.group(kind):
            raw = m.group(kind).replace(",", "")
            dp = len(raw.split(".")[1]) if "." in raw else 0
            return kind, float(raw), dp
    if m.group("day"):
        return "date", float(m.group("dyear")), 0
    return "year", float(m.group("year")), 0


def _matches(kind: str, value: float, dp: int, f: Fact) -> bool:
    tol = 0.5 * 10 ** -dp + 1e-9
    if kind == "eur":
        return f.kind in ("level", "change_eur") and (f.metric in METRIC and METRIC[f.metric].unit == "eur_m") \
            and abs(abs(f.value) - abs(value)) <= tol
    if kind == "pp":
        return f.kind == "change_pp" and abs(abs(f.value) - value) <= tol
    if kind == "pct":
        if f.kind == "change_pct":
            return abs(abs(f.value) - abs(value)) <= tol
        if f.kind == "const":
            return abs(f.value - value) <= tol
        return f.kind == "level" and f.metric in RATIOS and abs(f.value - value) <= tol
    if kind == "date":
        return f.kind == "date" and int(f.value) == int(value)
    return f.kind == "year" and int(f.value) == int(value)


def _context_score(f: Fact, before: str, after: str) -> float:
    """How well the words before a number, within its clause, fit a candidate fact (for numbers that several facts
    print alike)."""
    score = 0.0
    if f.metric in METRIC:
        near = before.lower()
        for s in METRIC[f.metric].synonyms + (METRIC[f.metric].label.lower(),):
            pos = near.rfind(s)
            if pos >= 0:
                # the closer and the more specific the label, the better
                score = max(score, 3.0 + pos / max(1, len(near)) + 0.02 * len(s))
    tail = before[-14:].lower()
    if f.kind == "const":
        # limits and the risk appetite are announced by their own words
        if re.search(r"(at least|range of|up to|below|floor of|cap of|limit of|appetite of)\s*$", before[-24:].lower()):
            score += 6.0
        elif re.search(r"\bto\s*$", tail) and re.search(r"range of\s*\S+\s+to\s*$", before[-30:].lower()):
            score += 6.0
    if f.period == "change":
        score += 2.0 if re.search(r"\bby\s*$", tail) else -1.0
    elif re.search(r"\bby\s*$", tail):
        score -= 2.0
    if f.period == "prev" and (re.search(r"\(\s*$", before[-3:]) or re.search(r"\bfrom\s*$", tail)):
        score += 1.5
    if f.period == "cur" and re.search(r"\bto\s*$", tail):
        score += 1.0
    if f.period == "prev" and re.search(r"\(\d{4}:\s*$", before[-8:]):
        score += 2.5
    return score


def link_numbers(text: str, facts: dict[str, Fact]) -> list[dict]:
    """Every number in a text with the facts it may stand for, the best one first."""
    out = []
    for m in NUMBER.finditer(text):
        kind, value, dp = _value_of(m)
        cands = [f for f in facts.values() if _matches(kind, value, dp, f)]
        start = max(text.rfind(". ", 0, m.start()), text.rfind("; ", 0, m.start()), m.start() - 250) + 1
        before, after = text[max(0, start):m.start()], text[m.end():m.end() + 60]
        ranked = sorted(cands, key=lambda f: -_context_score(f, before, after))
        out.append({"start": m.start(), "end": m.end(), "text": m.group(0), "kind": kind, "value": value, "dp": dp,
                    "candidates": [f.id for f in ranked]})
    return out


def templatize(text: str, facts: dict[str, Fact]) -> tuple[str, list[dict]]:
    """A published text turned back into a template: every number that a fact of its year prints is replaced by
    that fact's placeholder. Numbers that no fact prints stay as they are and are returned as unlinked."""
    links = link_numbers(text, facts)
    parts, pos, unlinked = [], 0, []
    for ln in links:
        parts.append(text[pos:ln["start"]])
        if ln["candidates"]:
            parts.append("{{" + ln["candidates"][0] + "}}")
        else:
            parts.append(ln["text"])
            unlinked.append(ln)
        pos = ln["end"]
    parts.append(text[pos:])
    return "".join(parts), unlinked


# ---------------------------------------------------------------------------
# The old method: last year's text, figures updated
# ---------------------------------------------------------------------------

def copy_and_update(company: Company, year: int) -> dict:
    """Last year's published E.1 and E.2, with every figure that code can link to last year's catalogue replaced
    by the same item of this year's catalogue. Words, causes and anything code cannot link stay as they were."""
    last = year - 1
    if last not in company.narratives:
        raise ValueError(f"no published text for {last}")
    f_last, f_now = facts_for(company, last), facts_for(company, year)
    slots, templates, changes = {}, {}, []
    for row in company.narratives[last]:
        template, unlinked = templatize(row["text"], f_last)
        templates[row["slot"]] = template
        slots[row["slot"]] = render(template, f_now)
        for pid in PLACEHOLDER.findall(template):
            if pid in f_now and pid in f_last and f_now[pid].display != f_last[pid].display:
                changes.append({"slot": row["slot"], "placeholder": pid, "before": f_last[pid].display,
                                "after": f_now[pid].display})
        for u in unlinked:
            changes.append({"slot": row["slot"], "placeholder": "", "before": u["text"], "after": u["text"]})
    return {"slots": slots, "templates": templates, "changes": changes, "source_year": last}


# ---------------------------------------------------------------------------
# Slots of E.1 and E.2 and the Art. 297 elements they carry
# ---------------------------------------------------------------------------

SLOTS: list[tuple[str, str, str]] = [
    ("e1_policy", "E.1", "Objectives, policies and processes for managing own funds"),
    ("e1_structure", "E.1", "Structure, amount and quality of own funds by tier"),
    ("e1_eligibility", "E.1", "Eligible own funds to cover the SCR and the MCR"),
    ("e1_reconciliation", "E.1", "Equity in the financial statements and excess of assets over liabilities"),
    ("e2_amounts", "E.2", "Solvency Capital Requirement and Minimum Capital Requirement"),
    ("e2_methods", "E.2", "Calculation methods and inputs to the MCR"),
    ("e2_changes", "E.2", "Material changes over the reporting period"),
]
SLOT_IDS = [s[0] for s in SLOTS]
SLOT_TITLE = {s[0]: s[2] for s in SLOTS}
SLOT_SECTION = {s[0]: s[1] for s in SLOTS}
TABLE_AFTER = {"e1_structure": "Table E.1", "e2_amounts": "Table E.2"}

# Art. 297 of Delegated Regulation (EU) 2015/35: the elements of E.1 and E.2 this app looks for.
ELEMENTS: list[dict] = [
    {"id": "297(1)(a)", "what": "objectives, policies and processes for managing own funds, incl. the time horizon "
                                "for business planning", "slots": ["e1_policy"],
     "all": [r"objective|polic(?:y|ies)|process", r"business plan|planning|time horizon"]},
    {"id": "297(1)(b)", "what": "structure, amount and quality of own funds by tier, with the changes in each tier",
     "slots": ["e1_structure"], "all": [r"tier\s?1"], "table": "Table E.1"},
    {"id": "297(1)(c)", "what": "eligible own funds to cover the SCR, by tier", "slots": ["e1_eligibility"],
     "all": [r"eligible", r"solvency capital requirement|\bscr\b"], "table": "Table E.1"},
    {"id": "297(1)(d)", "what": "eligible basic own funds to cover the MCR, by tier", "slots": ["e1_eligibility"],
     "all": [r"eligible", r"minimum capital requirement|\bmcr\b"], "table": "Table E.1"},
    {"id": "297(1)(e)", "what": "differences between equity in the financial statements and the excess of assets over "
                                "liabilities", "slots": ["e1_reconciliation"],
     "all": [r"financial statements", r"excess of assets over liabilities"]},
    {"id": "297(1)(g)", "what": "ancillary own funds (or a statement that there are none)",
     "slots": ["e1_structure", "e1_eligibility"], "all": [r"ancillary own funds?"]},
    {"id": "297(1)(h)", "what": "items deducted from own funds and restrictions on their availability",
     "slots": ["e1_eligibility", "e1_reconciliation", "e1_structure"], "all": [r"deduct|restrict"]},
    {"id": "297(2)(a)", "what": "the SCR and the MCR at the end of the reporting period", "slots": ["e2_amounts"],
     "facts": ["scr_cur", "mcr_cur"]},
    {"id": "297(2)(b)", "what": "the SCR split by risk module", "slots": ["e2_amounts"], "table": "Table E.2",
     "all": []},
    {"id": "297(2)(c)", "what": "simplified calculations used (or none)", "slots": ["e2_methods"],
     "all": [r"simplif"]},
    {"id": "297(2)(d)", "what": "undertaking-specific parameters used (or none)", "slots": ["e2_methods"],
     "all": [r"undertaking-specific parameters?"]},
    {"id": "297(2)(f)", "what": "any capital add-on (or none)", "slots": ["e2_methods", "e2_amounts"],
     "all": [r"capital add-on"]},
    {"id": "297(2)(g)", "what": "inputs used to calculate the MCR", "slots": ["e2_methods"],
     "all": [r"minimum capital requirement|\bmcr\b", r"input|linear|technical provisions|premiums|capital at risk"]},
    {"id": "297(2)(h)", "what": "material changes to the SCR and the MCR and their reasons", "slots": ["e2_changes"],
     "all": [r"solvency capital requirement|\bscr\b", r"minimum capital requirement|\bmcr\b"]},
]


# ---------------------------------------------------------------------------
# Checks
# ---------------------------------------------------------------------------

@dataclass
class Finding:
    check: str
    severity: str        # "error", "warning" or "info"
    slot: str
    quote: str
    message: str

    def to_dict(self) -> dict:
        return {"check": self.check, "severity": self.severity, "slot": self.slot, "quote": self.quote,
                "message": self.message}


CHECKS = {
    "N-LINK": "every number belongs to a figure of the reporting year or of its driver bullets",
    "N-DATE": "every month and year (e.g. 'March 2024') matches the Company sheet or a driver bullet",
    "N-YEAR": "every year is the reporting year or the previous year",
    "T-DIR": "every movement word matches the direction of the change",
    "T-HOUSE": "every movement is described with the house phrase for its size",
    "R-PP": "a change of a ratio is given in percentage points, not per cent",
    "C-CAUSE": "every stated cause is traced to a driver bullet of the reporting year",
    "B-COVER": "every driver bullet of the reporting year is reflected or deliberately left out",
    "G-TERM": "the house terminology is used, and no term it says to avoid",
    "E-ELEM": "the elements of Art. 297(1) and (2) are there",
}

# Movement verbs only: adjectives such as "higher premiums" or phrases such as "eligible up to" say no direction.
UP = r"increased|rose|risen|grew|went up|raised"
DOWN = r"decreased|fell|fallen|declined|dropped|went down|reduced|lowered"
STABLE = r"remained broadly stable|remained stable|remained unchanged|(?:was|were) unchanged|remained flat"
DIRECTION = re.compile(rf"\b(?P<up>{UP})\b|\b(?P<down>{DOWN})\b|\b(?P<stable>{STABLE})\b", re.IGNORECASE)
TRANSITIVE = re.compile(r"increased|raised|reduced|lowered|decreased", re.IGNORECASE)
INTENSITY = re.compile(r"\b(slightly|significantly|sharply|strongly|considerably|substantially|marginally|"
                       r"materially|somewhat|moderately)\b", re.IGNORECASE)
CAUSAL = re.compile(r"\b(because|due to|driven by|as a result of|resulting from|reflect(?:s|ed|ing)?|mainly from|"
                    r"owing to|caused by|led to|thanks to|following|attributable to|on account of|arising from|"
                    r"result(?:ed)? from|explained by)\b", re.IGNORECASE)
STOP = set("""a an the and or of to in on for by with from at as is was were be been are this that these those it its
our we us their which who whom whose than then into over under after before during also more less most least very
year years end period during while where when has have had not no nor any all each per since about between both such
other others there here same within without upon because due driven resulting reflect reflects reflected
reflecting mainly owing caused led thanks following attributable arising explained one two can may will
would could should did does done new""".split())


def split_sentences(text: str) -> list[str]:
    parts = re.split(r"(?<=[.!?])\s+(?=[A-Z(\"“])", " ".join(text.split()))
    return [p for p in parts if p.strip()]


def _clauses(sentence: str) -> list[str]:
    return [c for c in re.split(r";|,\s(?:while|whereas|and|but)\s|\swhile\s|\swhereas\s", sentence) if c.strip()]


def _metric_mentions(text: str) -> list[tuple[int, int, str]]:
    """(start, end, metric id) of every metric named in a text, longest names first, no overlaps."""
    low = text.lower()
    spans: list[tuple[int, int, str]] = []
    names = sorted(((s, m.id) for m in METRICS if m.movement for s in m.synonyms), key=lambda t: -len(t[0]))
    for name, mid in names:
        for mm in re.finditer(r"(?<![a-z])" + re.escape(name) + r"(?![a-z])", low):
            if not any(a < mm.end() and mm.start() < b for a, b, _ in spans):
                spans.append((mm.start(), mm.end(), mid))
    return sorted(spans)


def _words(text: str) -> set[str]:
    out = set()
    for w in re.findall(r"[a-z][a-z\-]{2,}", text.lower()):
        if w in STOP:
            continue
        for suffix in ("ing", "ed", "es", "s"):
            if w.endswith(suffix) and len(w) - len(suffix) >= 4:
                w = w[: -len(suffix)]
                break
        out.add(w)
    return out


def _movement_target(clause: str, mentions: list[tuple[int, int, str]], d: re.Match) -> str | None:
    """The metric a movement verb speaks about: the subject named just before it ("the SCR, which ..., fell"), or
    for a transitive verb the object right after it ("reduced Tier 1 own funds"). None if neither is clear."""
    before = [mm for mm in mentions if mm[1] <= d.start() and d.start() - mm[1] <= 140]
    if before:
        subject = before[-1]
        between = clause[subject[1]:d.start()]
        if not (CAUSAL.search(between) or DIRECTION.search(between)
                or re.search(r"\b(?:while|whereas|but|as|since|when|after|because)\b", between)):
            return subject[2]
    if TRANSITIVE.fullmatch(d.group(0)):
        after = [mm for mm in mentions if mm[0] >= d.end() and mm[0] - d.end() <= 18]
        if after:
            return after[0][2]
    return None


def _arithmetic_cause(sentence: str) -> bool:
    """True if the cause part of a sentence only relates figures to each other ("because eligible own funds rose
    faster than the SCR"): that is arithmetic, which the figures check, not a cause that needs a bullet."""
    m = CAUSAL.search(sentence)
    if not m:
        return False
    cause = sentence[m.end():].lower()
    for start, end, _ in reversed(_metric_mentions(cause)):
        cause = cause[:start] + " " + cause[end:]
    content = _words(DIRECTION.sub(" ", cause)) - {"faster", "slower", "more", "less", "than", "while", "both"}
    return len(content) <= 1


def word_weights(company: Company) -> dict[str, float]:
    """Inverse document frequency over all bullets of the workbook: words that every year's bullets use
    ('reserve', 'risk', 'own funds') weigh little, distinctive ones ('profit', 'annuities', 'hailstorm') much."""
    import math
    docs = [_words(b["text"] + " " + b["topic"]) for rows in company.bullets.values() for b in rows]
    df: dict[str, int] = {}
    for d in docs:
        for w in d:
            df[w] = df.get(w, 0) + 1
    n = max(1, len(docs))
    return {w: math.log(1 + n / k) for w, k in df.items()}


def _overlap(sentence: str, bullet: str, weights: dict[str, float] | None = None) -> float:
    a, b = _words(sentence), _words(bullet)
    if not weights:
        return len(a & b) / max(1, min(len(a), len(b)))
    top = max(weights.values(), default=1.0)

    def w(x: str) -> float:
        return weights.get(x, top)
    return sum(w(x) for x in a & b) / max(1e-9, min(sum(w(x) for x in a), sum(w(x) for x in b)))


def run_checks(company: Company, year: int, slots: dict[str, str], *, method: str = "") -> list[Finding]:
    """The deterministic checks on a rendered draft (no placeholders left)."""
    facts = facts_for(company, year)
    moves = movements(facts)
    findings: list[Finding] = []
    bullets_now = company.bullets.get(year, [])
    bullets_before = [(y, b) for y in sorted(company.bullets) if y < year for b in company.bullets[y]]
    bullet_numbers = {n.group(0) for b in bullets_now for n in NUMBER.finditer(b["text"])}
    weights = word_weights(company)
    month_re = re.compile(r"\b(" + "|".join(MONTHS) + r")(?:\s+(\d{4}))?\b")
    sources = [str(v) for v in company.text.values()] + [b["text"] for rows in company.bullets.values() for b in rows]
    dated = {(m.group(1), int(m.group(2))) for t in sources for m in month_re.finditer(t) if m.group(2)}
    undated_now = {m.group(1) for b in bullets_now for m in month_re.finditer(b["text"]) if not m.group(2)}
    older_facts = {y: facts_for(company, y) for y in company.years if y < year and y - 1 in company.figures}

    for slot, text in slots.items():
        for sentence in split_sentences(text):
            # ---- N-LINK and N-YEAR -------------------------------------------------------------------------
            for ln in link_numbers(sentence, facts):
                if ln["kind"] == "year":
                    month = re.search(r"\b(" + "|".join(MONTHS) + r")\s+$", sentence[max(0, ln["start"] - 12):ln["start"]])
                    if month and (month.group(1), int(ln["value"])) in dated:
                        continue  # a dated event that the Company sheet or a bullet confirms (see N-DATE)
                    if int(ln["value"]) not in (year, year - 1):
                        findings.append(Finding("N-YEAR", "warning", slot, sentence,
                                                f"{ln['text']} is neither {year} nor {year - 1}: check that the "
                                                "older year is meant"))
                    continue
                if ln["candidates"] or ln["text"] in bullet_numbers:
                    continue
                stale = [y for y, fy in older_facts.items()
                         if any(_matches(ln["kind"], ln["value"], ln["dp"], f) for f in fy.values())]
                if stale:
                    findings.append(Finding("N-LINK", "error", slot, sentence,
                                            f"{ln['text']} is not a {year} figure; it matches a figure of "
                                            f"{max(stale)}"))
                else:
                    findings.append(Finding("N-LINK", "error", slot, sentence,
                                            f"{ln['text']} matches no figure of {year} and no driver bullet"))
            rest = REFERENCE.sub(" ", NUMBER.sub(" ", sentence))
            for stray in STRAY_DIGITS.findall(rest):
                findings.append(Finding("N-LINK", "warning", slot, sentence,
                                        f"the number {stray} has no unit that the check can read; verify it by hand"))
            # ---- T-DIR, T-HOUSE and R-PP -----------------------------------------------------------------
            for clause in _clauses(sentence):
                mentions = _metric_mentions(clause)
                for d in DIRECTION.finditer(clause):
                    said = "increase" if d.group("up") else "decrease" if d.group("down") else "unchanged"
                    target = _movement_target(clause, mentions, d)
                    if target is None:
                        continue
                    mv = moves.get(target)
                    if mv is None:
                        continue
                    stable_house = bool(mv.phrase) and "stable" in mv.phrase
                    if mv.direction != said and not (said == "unchanged" and stable_house):
                        went = {"increase": "rose", "decrease": "fell"}.get(mv.direction, "did not change")
                        arrow = {"increase": " up", "decrease": " down"}.get(mv.direction, "")
                        findings.append(Finding("T-DIR", "error", slot, sentence,
                                                f"'{d.group(0)}' for the {METRIC[target].label}, which {went} "
                                                f"({mv.display}{arrow})"))
                        continue
                    # The house phrase is checked where a movement is quantified ("... by 2.5%") or called stable.
                    quantified = re.match(r"[a-z ]{0,25}?\bby\s+(?:EUR\s)?\d", clause[d.end():d.end() + 45],
                                          re.IGNORECASE)
                    if mv.phrase and (quantified or said == "unchanged"):
                        window = clause[max(0, d.start() - 5):d.end() + 25].lower()
                        if mv.phrase.lower() not in window:
                            findings.append(Finding("T-HOUSE", "warning", slot, sentence,
                                                    f"the house phrase for the {METRIC[target].label} "
                                                    f"({mv.display}) is '{mv.phrase}'"))
                for mm in mentions:
                    if mm[2] in RATIOS:
                        tail = clause[mm[1]:mm[1] + 90]
                        hit = re.search(r"\bby\s(\d+(?:\.\d+)?)\s?%", tail)
                        if hit:
                            findings.append(Finding("R-PP", "error", slot, sentence,
                                                    f"'by {hit.group(1)}%' for the {METRIC[mm[2]].label}: a change "
                                                    "of a ratio is given in percentage points"))
            # ---- C-CAUSE (deterministic part): is the cause traced to a bullet of this year? -------------
            marker = CAUSAL.search(sentence)
            if marker and not _arithmetic_cause(sentence):
                cause = sentence[marker.start():]
                probe = cause if len(_words(cause)) >= 2 else sentence
                now = max(((_overlap(probe, b["text"], weights), b) for b in bullets_now), default=(0.0, None),
                          key=lambda t: t[0])
                old = max(((_overlap(probe, b["text"], weights), y, b) for y, b in bullets_before),
                          default=(0.0, 0, None), key=lambda t: t[0])
                if old[0] >= 0.34 and old[0] > now[0] + 0.25:
                    findings.append(Finding("C-CAUSE", "warning", slot, sentence,
                                            f"the cause resembles {old[1]} bullet {old[2]['id']} ('{old[2]['topic']}'), "
                                            f"not a {year} bullet: a cause carried over from an earlier year?"))
                elif now[0] < 0.34:
                    findings.append(Finding("C-CAUSE", "warning", slot, sentence,
                                            "the cause could not be traced to a driver bullet of "
                                            f"{year} (not verifiable by code)"))

    # ---- N-DATE --------------------------------------------------------------------------------------------
    for slot, text in slots.items():
        for sentence in split_sentences(text):
            for m in month_re.finditer(sentence):
                if not m.group(2) or re.search(r"\b\d{1,2}\s$", sentence[max(0, m.start() - 4):m.start()]):
                    continue  # no year, or a full date such as 31 December 2026 (checked by N-LINK)
                month, yr = m.group(1), int(m.group(2))
                if (month, yr) in dated or (month in undated_now and yr == year):
                    continue
                other = sorted(y for mo, y in dated if mo == month)
                if other:
                    findings.append(Finding("N-DATE", "error", slot, sentence,
                                            f"'{month} {yr}': the Company sheet and the bullets say "
                                            + ", ".join(f"{month} {y}" for y in other)))
                else:
                    findings.append(Finding("N-DATE", "warning", slot, sentence,
                                            f"'{month} {yr}' is in no driver bullet and not in the Company sheet: "
                                            "verify"))

    # ---- B-COVER -------------------------------------------------------------------------------------------
    sentences = [s for t in slots.values() for s in split_sentences(t)]
    for b in bullets_now:
        best = max((_overlap(s, b["text"], weights) for s in sentences), default=0.0)
        if best < 0.3:
            findings.append(Finding("B-COVER", "info", "", b["text"],
                                    f"bullet {b['id']} ('{b['topic']}') is not reflected in the text: "
                                    "left out on purpose?"))

    # ---- G-TERM --------------------------------------------------------------------------------------------
    for slot, text in slots.items():
        for term in company.terminology:
            for avoid in term["avoid"]:
                for mm in re.finditer(r"(?<![A-Za-z])" + re.escape(avoid) + r"(?![A-Za-z])", text, re.IGNORECASE):
                    sentence = next((s for s in split_sentences(text) if mm.group(0) in s), mm.group(0))
                    findings.append(Finding("G-TERM", "warning", slot, sentence,
                                            f"'{mm.group(0)}': the house term is '{term['preferred']}'"))

    # ---- E-ELEM --------------------------------------------------------------------------------------------
    for el in ELEMENTS:
        text = " ".join(slots.get(s, "") for s in el["slots"]).lower()
        ok = all(re.search(p, text) for p in el.get("all", []))
        if el.get("facts"):
            ok = ok and all(facts[fid].display.lower() in text for fid in el["facts"] if fid in facts)
        if not ok and el.get("table") and not el.get("all") and not el.get("facts"):
            ok = True
        if not ok:
            findings.append(Finding("E-ELEM", "warning", ", ".join(el["slots"]), "",
                                    f"Art. {el['id']}: {el['what']} - not found in the text"))
    return findings


def element_status(company: Company, year: int, slots: dict[str, str]) -> list[dict]:
    """One row per Art. 297 element: found in the text, inserted by code as a table, or not found."""
    facts = facts_for(company, year)
    rows = []
    for el in ELEMENTS:
        text = " ".join(slots.get(s, "") for s in el["slots"]).lower()
        found = all(re.search(p, text) for p in el.get("all", []))
        if el.get("facts"):
            found = found and all(facts[f].display.lower() in text for f in el["facts"] if f in facts)
        where = "in the text" if found and (el.get("all") or el.get("facts")) else ""
        if el.get("table"):
            where = (where + " and " if where else "") + f"in {el['table']} (inserted by code)"
            found = True if not (el.get("all") or el.get("facts")) else found
        rows.append({"element": f"Art. {el['id']}", "what": el["what"], "status": "found" if found else "not found",
                     "where": where or "-"})
    return rows


def similarity(a: str, b: str) -> float:
    """Word-level similarity of two texts (0 to 1), for comparing a draft's wording with last year's."""
    return difflib.SequenceMatcher(None, a.split(), b.split(), autojunk=False).ratio()


# ---------------------------------------------------------------------------
# The two tables that code inserts
# ---------------------------------------------------------------------------

def table_e1(company: Company, year: int) -> list[dict]:
    rows = []
    for tier, lab in (("t1", "Tier 1"), ("t2", "Tier 2"), ("t3", "Tier 3")):
        cur, prev = company.figures[year], company.figures[year - 1]
        if abs(cur[f"of_{tier}"]) + abs(prev[f"of_{tier}"]) < 1e-9:
            continue
        rows.append({
            "": lab,
            f"available {year}": cur[f"of_{tier}"],
            f"eligible for the SCR {year}": cur.get(f"eof_scr_{tier}", 0.0),
            f"eligible for the MCR {year}": cur.get(f"eof_mcr_{tier}", 0.0),
            f"available {year - 1}": prev[f"of_{tier}"],
            f"eligible for the SCR {year - 1}": prev.get(f"eof_scr_{tier}", 0.0),
            f"eligible for the MCR {year - 1}": prev.get(f"eof_mcr_{tier}", 0.0),
        })
    cur, prev = company.figures[year], company.figures[year - 1]
    rows.append({"": "Total", f"available {year}": cur["of_total"], f"eligible for the SCR {year}": cur["eof_scr"],
                 f"eligible for the MCR {year}": cur["eof_mcr"], f"available {year - 1}": prev["of_total"],
                 f"eligible for the SCR {year - 1}": prev["eof_scr"], f"eligible for the MCR {year - 1}": prev["eof_mcr"]})
    return rows


def table_e2(company: Company, year: int) -> list[dict]:
    cur, prev = company.figures[year], company.figures[year - 1]
    spec = [("scr_market", "Market risk", 1), ("scr_default", "Counterparty default risk", 1),
            ("scr_life", "Life underwriting risk", 1), ("scr_health", "Health underwriting risk", 1),
            ("scr_nonlife", "Non-life underwriting risk", 1), ("div_benefit", "Diversification", -1),
            ("bscr", "Basic Solvency Capital Requirement", 1), ("scr_operational", "Operational risk", 1),
            ("lac_tp", "Loss-absorbing capacity of technical provisions", -1),
            ("lac_dt", "Loss-absorbing capacity of deferred taxes", -1),
            ("scr", "Solvency Capital Requirement (SCR)", 1), ("mcr", "Minimum Capital Requirement (MCR)", 1),
            ("eof_scr", "Eligible own funds to cover the SCR", 1), ("cov_scr", "SCR coverage ratio", 1),
            ("eof_mcr", "Eligible own funds to cover the MCR", 1), ("cov_mcr", "MCR coverage ratio", 1)]
    rows = []
    for mid, lab, sign in spec:
        if abs(cur[mid]) + abs(prev[mid]) < 1e-9:
            continue
        if mid in RATIOS:
            rows.append({"": lab, str(year): fmt_pct(cur[mid]), str(year - 1): fmt_pct(prev[mid])})
        else:
            rows.append({"": lab, str(year): f"{sign * cur[mid]:,.1f}", str(year - 1): f"{sign * prev[mid]:,.1f}"})
    return rows
